import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HEADERS = ["공정", "위험요소", "위험등급", "근거 사고사례", "저감대책", "2차 위험요소"]

FIELD_TO_HEADER = {
    "위험요소": "위험요소",
    "위험등급": "위험등급",
    "근거 사고사례": "근거 사고사례",
    "저감대책": "저감대책",
    "저감대책 적용 후 2차 위험요소": "2차 위험요소",
}


def parse_report_to_rows(report: str) -> list[dict]:
    """에이전트가 생성한 마크다운 보고서를 표 형태(행 리스트)로 파싱한다."""
    sections = re.split(r"\n(?=### )", report)
    rows = []

    for section in sections:
        header_match = re.match(r"### (.+)", section.strip())
        if not header_match:
            continue
        process = header_match.group(1).strip()
        if "제외된 후보" in process:
            continue

        row = {h: "" for h in HEADERS}
        row["공정"] = process
        for field, header in FIELD_TO_HEADER.items():
            m = re.search(rf"\*\*{re.escape(field)}\*\*:\s*(.+)", section)
            if m:
                row[header] = m.group(1).strip()
        rows.append(row)

    return rows


def export_to_excel(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "위험성평가표"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])

    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {"공정": 20, "위험요소": 30, "위험등급": 15, "근거 사고사례": 25, "저감대책": 35, "2차 위험요소": 35}
    for i, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[chr(64 + i)].width = widths.get(header, 20)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
